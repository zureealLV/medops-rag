"""Build deterministic medical-device files for manual multimodal ingestion tests."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "sample_data" / "multimodal_demo"
FONT_REGULAR = Path("C:/Windows/Fonts/msyh.ttc")
FONT_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")


def _font(size: int, *, bold: bool = False):
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size)


def _rounded_panel(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], fill: str) -> None:
    draw.rounded_rectangle(box, radius=28, fill=fill, outline="#9bc9c4", width=3)


def build_raster_assets() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)

    pump = Image.new("RGB", (1400, 900), "#edf7f5")
    draw = ImageDraw.Draw(pump)
    _rounded_panel(draw, (90, 70, 1310, 830), "#ffffff")
    draw.text((150, 120), "MEDOPS DEMO INFUSION PUMP", font=_font(46, bold=True), fill="#164b4a")
    draw.rounded_rectangle((150, 220, 1250, 590), radius=24, fill="#193f45")
    draw.text((210, 275), "UPSTREAM OCCLUSION", font=_font(62, bold=True), fill="#ffcf6e")
    draw.text((210, 380), "INFUSION PAUSED", font=_font(54, bold=True), fill="#ffffff")
    draw.text((210, 485), "DEVICE ID  DEMO-PUMP-042", font=_font(34), fill="#aee7df")
    draw.text((150, 665), "TRAINING FIXTURE - NOT A REAL DEVICE", font=_font(34, bold=True), fill="#b44754")
    draw.text(
        (150, 725),
        "Follow the manufacturer instructions and facility protocol.",
        font=_font(29),
        fill="#567577",
    )
    pump.save(OUTPUT / "infusion_pump_alarm_panel.png", optimize=True)

    oximeter = Image.new("RGB", (1200, 800), "#f5fbfa")
    draw = ImageDraw.Draw(oximeter)
    _rounded_panel(draw, (80, 60, 1120, 740), "#ffffff")
    draw.text((135, 115), "PULSE OXIMETER REFERENCE", font=_font(44, bold=True), fill="#164b4a")
    draw.rounded_rectangle((135, 220, 1065, 585), radius=28, fill="#102f38")
    draw.text((205, 285), "SpO2", font=_font(40), fill="#aee7df")
    draw.text((205, 345), "98 %", font=_font(84, bold=True), fill="#68e0ce")
    draw.text((650, 285), "PULSE", font=_font(40), fill="#aee7df")
    draw.text((650, 345), "72 bpm", font=_font(78, bold=True), fill="#ffd173")
    draw.text(
        (135, 635), "SIMULATED DISPLAY - EDUCATIONAL TEST ONLY", font=_font(30, bold=True), fill="#b44754"
    )
    oximeter.save(OUTPUT / "pulse_oximeter_display.jpg", quality=94, optimize=True)
    oximeter.save(OUTPUT / "pulse_oximeter_display.webp", quality=92, method=6)

    rows = [
        {
            "asset_id": "DEMO-PUMP-042",
            "device_type": "infusion_pump",
            "department": "simulation_lab",
            "status": "inspection_due",
            "patient_data": "none",
        },
        {
            "asset_id": "DEMO-SPO2-017",
            "device_type": "pulse_oximeter",
            "department": "simulation_lab",
            "status": "ready",
            "patient_data": "none",
        },
    ]
    with (OUTPUT / "medical_device_inventory.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    with (OUTPUT / "medical_device_events.jsonl").open("w", encoding="utf-8", newline="\n") as stream:
        for item in (
            {
                "event_id": "DEMO-EVT-001",
                "asset_id": "DEMO-PUMP-042",
                "event": "upstream_occlusion_alarm",
                "action": "removed_from_service_for_authorized_review",
            },
            {
                "event_id": "DEMO-EVT-002",
                "asset_id": "DEMO-SPO2-017",
                "event": "functional_check",
                "action": "passed_in_simulation_lab",
            },
        ):
            stream.write(json.dumps(item, ensure_ascii=False) + "\n")


def build_docx() -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    OUTPUT.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    title = document.add_heading("输液泵报警教学测试资料", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = "Microsoft YaHei"
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.color.rgb = RGBColor(24, 54, 58)
    intro = document.add_paragraph(
        "本文件用于验证 DOCX 文本、表格、内嵌图片与 OCR 摄取。内容完全合成，不对应真实设备或患者。"
    )
    intro.style = document.styles["Normal"]
    document.add_heading("报警信息", level=1)
    document.add_paragraph(
        "模拟设备显示 UPSTREAM OCCLUSION 和 INFUSION PAUSED。真实设备应按照制造商说明与机构流程处置。"
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("字段", "模拟值", "用途")
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in (
        ("Device ID", "DEMO-PUMP-042", "关联合成资产"),
        ("Alarm", "UPSTREAM OCCLUSION", "验证表格检索"),
        ("Status", "INFUSION PAUSED", "验证状态文本"),
    ):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    document.add_heading("内嵌视觉证据", level=1)
    document.add_picture(str(OUTPUT / "infusion_pump_alarm_panel.png"), width=Cm(15.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "来源依据：FDA Infusion Pumps 公开说明；本文件为重新编写的教学测试样例，不包含真实医疗数据。"
    )
    document.save(OUTPUT / "infusion_pump_training_pack.docx")


def build_pdf() -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image as PdfImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    OUTPUT.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont("MicrosoftYaHei", str(FONT_REGULAR), subfontIndex=0))
    pdfmetrics.registerFont(TTFont("MicrosoftYaHeiBold", str(FONT_BOLD), subfontIndex=0))
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TitleCN",
        parent=styles["Title"],
        fontName="MicrosoftYaHeiBold",
        fontSize=22,
        leading=30,
        textColor=HexColor("#18363a"),
        spaceAfter=18,
    )
    heading = ParagraphStyle(
        "HeadingCN",
        parent=styles["Heading2"],
        fontName="MicrosoftYaHeiBold",
        fontSize=14,
        leading=21,
        textColor=HexColor("#137a74"),
        spaceBefore=12,
        spaceAfter=8,
    )
    body = ParagraphStyle(
        "BodyCN",
        parent=styles["BodyText"],
        fontName="MicrosoftYaHei",
        fontSize=10.5,
        leading=18,
        textColor=HexColor("#314d50"),
        spaceAfter=8,
    )
    path = OUTPUT / "medical_device_quick_reference.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=48,
        bottomMargin=48,
        title="医疗器械多模态摄取测试资料",
    )
    story = [
        Paragraph("医疗器械多模态摄取测试资料", title),
        Paragraph(
            "用于验证 PDF 文本抽取、表格、页码和视觉证据。资料为合成教学内容，不对应真实患者或真实设备。",
            body,
        ),
        Paragraph("脉搏血氧仪", heading),
        Paragraph(
            "脉搏血氧仪通过光学方式估算血氧饱和度和脉率。末梢循环、皮肤与指甲状况、运动、传感器位置和设备状态都可能影响读数。",
            body,
        ),
        PdfImage(str(OUTPUT / "pulse_oximeter_display.jpg"), width=470, height=313),
        Spacer(1, 12),
        Paragraph("输液泵报警测试", heading),
    ]
    data = [
        ["Device ID", "Alarm", "Expected handling"],
        ["DEMO-PUMP-042", "UPSTREAM OCCLUSION", "Follow authorized protocol"],
        ["DEMO-SPO2-017", "FUNCTIONAL CHECK", "Simulation result only"],
    ]
    table = Table(data, colWidths=[115, 145, 210])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "MicrosoftYaHei"),
                ("FONTNAME", (0, 0), (-1, 0), "MicrosoftYaHeiBold"),
                ("BACKGROUND", (0, 0), (-1, 0), HexColor("#dff4ef")),
                ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#18363a")),
                ("GRID", (0, 0), (-1, -1), 0.6, HexColor("#b9cfcc")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend(
        [
            table,
            Spacer(1, 14),
            Paragraph(
                "来源依据：FDA Pulse Oximeters 与 Infusion Pumps 公开页面。"
                "使用前应核对制造商说明和机构制度。",
                body,
            ),
        ]
    )
    doc.build(story)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=("raster", "docx", "pdf", "all"))
    args = parser.parse_args()
    if args.mode in {"raster", "all"}:
        build_raster_assets()
    if args.mode in {"docx", "all"}:
        build_docx()
    if args.mode in {"pdf", "all"}:
        build_pdf()


if __name__ == "__main__":
    main()
